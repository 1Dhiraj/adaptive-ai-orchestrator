"""Build the base-paper study report as a Word document.

Everything here is either (a) bibliographic metadata read from the paper's
IEEE Xplore record, (b) the abstract quoted with attribution, or (c) our own
analysis of how the paper relates to this project. The full text is paywalled
and has not been read, and the document says so on its face.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUTPUT = Path("Base Paper Study - Multi-Agent Orchestrator System.docx")

NAVY = RGBColor(0x1F, 0x3B, 0x57)
DARK = RGBColor(0x21, 0x25, 0x29)
GREY = RGBColor(0x55, 0x5B, 0x61)

PAPER = {
    "title": "Dynamic Workflow Orchestration with Executability Verification for "
             "Multi-Agent AutoML Task Execution",
    "authors": "Luke Deng, Jie Yan, Cheng Cao, Mingyang Zhao, Songchang Jin",
    "venue": "2026 6th International Conference on Artificial Intelligence, "
             "Big Data and Algorithms (CAIBDA)",
    "publisher": "IEEE",
    "date": "12 June 2026",
    "doi": "10.1109/CAIBDA70336.2026.11621590",
    "url": "https://ieeexplore.ieee.org/document/11621590",
    "indexing": "IEEE Xplore",
}

# Quoted with attribution from the publicly displayed IEEE Xplore record.
ABSTRACT = (
    "Automated machine learning (AutoML) engineering requests often involve multiple "
    "interdependent stages, heterogeneous tools, and strict execution constraints, making "
    "them difficult for existing large language model (LLM)-based agents to plan and execute "
    "reliably. In this paper, we propose a workflow-orchestrated multi-agent framework that "
    "explicitly models AutoML task execution as a directed acyclic graph (DAG). A workflow "
    "orchestrator (WO) parses user requirements, constructs or retrieves a task graph, "
    "dynamically assigns executable subtasks to specialized agents, and coordinates execution "
    "under runtime executability verification. To improve robustness, each subagent checks "
    "input availability, capability–task matching, and constraint compliance before "
    "execution, and returns structured diagnostic feedback when failures occur. WO aggregates "
    "intermediate results, evaluates the entire workflow using a transparent global acceptance "
    "score, and performs targeted replanning through a closed-loop repair mechanism. We "
    "evaluate the proposed framework on a task-level benchmark covering typical ML workflow "
    "stages, including data generation, data annotation, model recommendation, training, "
    "evaluation, deployment, and abnormal requests. Results show that the proposed method "
    "improves workflow-level success and recovery performance under fixed planning budgets "
    "compared with planning-driven and hierarchical orchestration baselines."
)

KEYWORDS = ("Large language models; Multi-agent systems; Automated machine learning; "
            "Planning; Task execution; Benchmark")


def shade(cell, hex_colour: str) -> None:
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hex_colour)
    cell._tc.get_or_add_tcPr().append(el)


def style_doc(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = DARK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(0.8)
        section.left_margin = section.right_margin = Inches(0.9)


def heading(doc: Document, text: str, size: int = 13) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = NAVY


def bullet(doc: Document, text: str) -> None:
    """A bullet whose '**bold**' spans become real bold runs."""
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    for i, chunk in enumerate(text.split("**")):
        if chunk:
            run = p.add_run(chunk)
            run.bold = (i % 2 == 1)


def kv_table(doc: Document, rows) -> None:
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for key, val in rows:
        cells = t.add_row().cells
        cells[0].width = Inches(1.5)
        cells[1].width = Inches(5.2)
        kr = cells[0].paragraphs[0].add_run(key)
        kr.bold = True
        kr.font.size = Pt(10.5)
        vr = cells[1].paragraphs[0].add_run(val)
        vr.font.size = Pt(10.5)
        shade(cells[0], "EEF2F6")


def build() -> None:
    doc = Document()
    style_doc(doc)

    # -- cover block ------------------------------------------------------
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("BASE PAPER STUDY REPORT")
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = NAVY

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("Adaptive Multi-Agent Orchestrator System")
    sr.bold = True
    sr.font.size = Pt(13)
    sr.font.color.rgb = DARK

    for line in [
        "J A L Dhiraj [913123104038]    •    N Nithersan [913123104108]",
        "Supervisor: Mrs. S. Sahebzathi",
        "Department of Computer Science and Engineering, VCET, Madurai",
        "First Review — 12/08/2026",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(line)
        run.font.size = Pt(10.5)
        run.font.color.rgb = GREY

    # -- 1. paper details --------------------------------------------------
    heading(doc, "1.  Base Paper Details")
    kv_table(doc, [
        ("Title", PAPER["title"]),
        ("Authors", PAPER["authors"]),
        ("Published in", f'{PAPER["venue"]}, {PAPER["publisher"]}'),
        ("Date", PAPER["date"]),
        ("DOI", PAPER["doi"]),
        ("Link", PAPER["url"]),
        ("Indexed in", PAPER["indexing"]),
        ("Keywords", KEYWORDS),
    ])

    # -- 2. abstract -------------------------------------------------------
    heading(doc, "2.  Abstract (as published)")
    note = doc.add_paragraph()
    nr = note.add_run("Reproduced from the paper's IEEE Xplore record and attributed to its "
                      "authors. Quoted here for study purposes only.")
    nr.italic = True
    nr.font.size = Pt(9.5)
    nr.font.color.rgb = GREY

    q = doc.add_paragraph()
    q.paragraph_format.left_indent = Inches(0.3)
    q.paragraph_format.right_indent = Inches(0.3)
    qr = q.add_run(f'“{ABSTRACT}”')
    qr.font.size = Pt(10.5)
    qr.italic = True

    # -- 3. approach -------------------------------------------------------
    heading(doc, "3.  Summary of the Proposed Approach")
    doc.add_paragraph(
        "The paper addresses a problem that appears whenever an AutoML request spans several "
        "dependent stages: an LLM agent must plan across interdependent steps, pick among "
        "heterogeneous tools, and respect execution constraints. The authors argue that "
        "existing LLM-based agents plan and execute such requests unreliably, and propose a "
        "workflow-orchestrated multi-agent framework in response.")
    bullet(doc, "**Explicit DAG model.** AutoML task execution is modelled as a directed "
                "acyclic graph rather than a linear chain, so dependencies between stages are "
                "represented directly.")
    bullet(doc, "**Workflow Orchestrator (WO).** A central component parses the user "
                "requirement, constructs or retrieves a task graph, dynamically assigns "
                "executable subtasks to specialised agents, and coordinates their execution.")
    bullet(doc, "**Runtime executability verification.** Before executing, each subagent "
                "checks input availability, capability–task matching, and constraint "
                "compliance, and returns structured diagnostic feedback when a failure occurs.")
    bullet(doc, "**Global acceptance score.** The orchestrator aggregates intermediate "
                "results and evaluates the whole workflow with a transparent score rather "
                "than judging each step in isolation.")
    bullet(doc, "**Closed-loop repair.** Failures trigger targeted replanning instead of "
                "restarting the workflow from the beginning.")

    heading(doc, "4.  Evaluation Reported by the Authors")
    doc.add_paragraph(
        "The framework is evaluated on a task-level benchmark covering typical machine "
        "learning workflow stages — data generation, data annotation, model "
        "recommendation, training, evaluation, deployment, and abnormal requests. The authors "
        "report improved workflow-level success and recovery performance under fixed planning "
        "budgets, compared against planning-driven and hierarchical orchestration baselines.")

    # -- 5. relevance ------------------------------------------------------
    heading(doc, "5.  Relevance to Our Project")
    doc.add_paragraph(
        "This paper is the closest published work to our system. It shares our structural "
        "assumptions, which makes it a suitable base paper, and the specific thing it does "
        "not do defines the space our contribution occupies.")

    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    for i, h in enumerate(["Aspect", "Base paper", "Our system"]):
        run = hdr[i].paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(hdr[i], "1F3B57")

    for aspect, base, ours in [
        ("Task representation",
         "Task execution modelled explicitly as a DAG",
         "Same — our Dependency Graph is a validated DAG with topological levels"),
        ("Orchestration",
         "A Workflow Orchestrator builds the task graph and assigns subtasks to "
         "specialised agents",
         "Same structure — our Task Planner builds the graph and the Agent Manager "
         "creates role-specialised agents per task"),
        ("Tool / capability handling",
         "Each subagent verifies capability–task matching before executing",
         "Generalised into capability-based tool routing, with automatic fallback to a "
         "capability-equivalent tool when one is unavailable"),
        ("Response to failure",
         "Structured diagnostics, then targeted replanning through a closed-loop repair "
         "mechanism",
         "Automatic substitution of an equivalent tool, so execution continues without "
         "replanning where possible"),
        ("Response to a requirement change",
         "Not addressed — the framework repairs after failures, not after a change "
         "in what the user asked for",
         "Input fingerprinting decides which steps a change actually invalidates; "
         "unaffected steps are reused at zero token cost"),
        ("Irreversible actions",
         "Not addressed",
         "Human approval gate holds the generated payload, so sign-off costs no "
         "additional LLM call"),
    ]:
        cells = t.add_row().cells
        for i, text in enumerate((aspect, base, ours)):
            run = cells[i].paragraphs[0].add_run(text)
            run.font.size = Pt(10)
            if i == 0:
                run.bold = True
                shade(cells[i], "EEF2F6")

    # -- 6. gap ------------------------------------------------------------
    heading(doc, "6.  Research Gap Identified from the Base Paper")
    doc.add_paragraph(
        "The base paper makes a workflow robust against failure: if a step fails, it is "
        "diagnosed and the plan is repaired. It does not address the separate question of "
        "what remains valid when the user changes the requirement itself. In that case a "
        "failure never occurs — every step succeeded — yet some results are now "
        "stale and others are still perfectly good.")
    bullet(doc, "There is no mechanism to determine **which steps a requirement change "
                "actually invalidates**, so adaptation defaults to re-running work that was "
                "still valid.")
    bullet(doc, "Verification happens **before** a step runs. Nothing compares a step's "
                "current inputs against the inputs it was last executed with.")
    bullet(doc, "Recovery depends on **replanning**; there is no capability-equivalent "
                "substitution that would let execution continue without a new plan.")

    heading(doc, "7.  How Our Work Extends It")
    doc.add_paragraph(
        "We keep the DAG model and the orchestrator-assigns-to-specialists structure, and add "
        "a decision procedure for reuse. Each step is hashed over its own definition together "
        "with the exact outputs of its dependencies. A step is then re-executed only if it is "
        "forced, has no cached result, or its fingerprint no longer matches.")
    doc.add_paragraph(
        "Because the hash covers dependency outputs, a change propagates by itself — and "
        "stops as soon as a re-run reproduces identical output. Measured against a restart-all "
        "baseline and LangGraph over six scenarios, this reduced step executions by 24% and "
        "token consumption by 22%. The saving depends on where the change lands: it is 0% when "
        "a change invalidates the entire graph, and up to 40% for late or no-op changes.")

    # -- 8. reference ------------------------------------------------------
    heading(doc, "8.  Reference")
    p = doc.add_paragraph()
    r = p.add_run(
        f'[1] L. Deng, J. Yan, C. Cao, M. Zhao and S. Jin, “{PAPER["title"]},” '
        f'in {PAPER["venue"]}, {PAPER["publisher"]}, 2026. '
        f'doi: {PAPER["doi"]}.')
    r.font.size = Pt(10.5)

    # -- sourcing note -----------------------------------------------------
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run(
        "Note on sources: the bibliographic details and the abstract above were taken from "
        "the paper's IEEE Xplore record. The full text is behind IEEE's paywall and was not "
        "consulted in preparing this report, so Sections 3 and 4 describe only what the "
        "published abstract states. Sections 5 to 7 are our own analysis. The full paper "
        "should be obtained through the college's IEEE subscription before the comparison in "
        "Section 5 is relied on in the final report.")
    r.italic = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = GREY

    doc.save(str(OUTPUT))
    print(f"written: {OUTPUT.resolve()}")


if __name__ == "__main__":
    build()
