"""One-page reference sheet for the base paper: citation, abstract, keywords.

Paper details only -- no student or project content. Everything is taken from
the paper's IEEE Xplore record; the abstract is quoted and attributed. This is
a reference sheet about the paper, not a copy of it.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUTPUT = Path("Base Paper - Reference Sheet.docx")

NAVY = RGBColor(0x1F, 0x3B, 0x57)
DARK = RGBColor(0x21, 0x25, 0x29)
GREY = RGBColor(0x55, 0x5B, 0x61)

TITLE = ("Dynamic Workflow Orchestration with Executability Verification "
         "for Multi-Agent AutoML Task Execution")
AUTHORS = "Luke Deng, Jie Yan, Cheng Cao, Mingyang Zhao, Songchang Jin"
VENUE = ("2026 6th International Conference on Artificial Intelligence, "
         "Big Data and Algorithms (CAIBDA)")
DATE = "12 June 2026"
DOI = "10.1109/CAIBDA70336.2026.11621590"
URL = "https://ieeexplore.ieee.org/document/11621590"

ABSTRACT = (
    "Automated machine learning (AutoML) engineering requests often involve multiple "
    "interdependent stages, heterogeneous tools, and strict execution constraints, making "
    "them difficult for existing large language model (LLM)-based agents to plan and execute "
    "reliably. In this paper, we propose a workflow-orchestrated multi-agent framework that "
    "explicitly models AutoML task execution as a directed acyclic graph (DAG). A workflow "
    "orchestrator (WO) parses user requirements, constructs or retrieves a task graph, "
    "dynamically assigns executable subtasks to specialized agents, and coordinates execution "
    "under runtime executability verification. To improve robustness, each subagent checks "
    "input availability, capability–task matching, and constraint compliance before "
    "execution, and returns structured diagnostic feedback when failures occur. WO aggregates "
    "intermediate results, evaluates the entire workflow using a transparent global acceptance "
    "score, and performs targeted replanning through a closed-loop repair mechanism. We "
    "evaluate the proposed framework on a task-level benchmark covering typical ML workflow "
    "stages, including data generation, data annotation, model recommendation, training, "
    "evaluation, deployment, and abnormal requests. Results show that the proposed method "
    "improves workflow-level success and recovery performance under fixed planning budgets "
    "compared with planning-driven and hierarchical orchestration baselines."
)

KEYWORDS = ("Large language models; Multi-agent systems; Automated machine learning; "
            "Planning; Task execution; Modeling; Training; Benchmark")


def shade(cell, hex_colour: str) -> None:
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hex_colour)
    cell._tc.get_or_add_tcPr().append(el)


def build() -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.font.color.rgb = DARK
    normal.paragraph_format.space_after = Pt(6)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.9)
        s.left_margin = s.right_margin = Inches(1.0)

    # -- title -----------------------------------------------------------
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(TITLE)
    r.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = NAVY

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(AUTHORS)
    r.font.size = Pt(11.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run(f"{VENUE}\nIEEE, {DATE}")
    r.italic = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = GREY

    # -- publication details ---------------------------------------------
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for key, val in [
        ("Title", TITLE),
        ("Authors", AUTHORS),
        ("Published in", VENUE),
        ("Publisher", "IEEE"),
        ("Date of publication", DATE),
        ("DOI", DOI),
        ("IEEE Xplore link", URL),
        ("Index terms", KEYWORDS),
    ]:
        cells = t.add_row().cells
        cells[0].width = Inches(1.4)
        cells[1].width = Inches(5.1)
        kr = cells[0].paragraphs[0].add_run(key)
        kr.bold = True
        kr.font.size = Pt(10)
        vr = cells[1].paragraphs[0].add_run(val)
        vr.font.size = Pt(10)
        shade(cells[0], "EEF2F6")

    # -- abstract ---------------------------------------------------------
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Abstract")
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = NAVY

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(ABSTRACT)
    r.font.size = Pt(10.5)

    # -- citation ---------------------------------------------------------
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Citation (IEEE format)")
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = NAVY

    p = doc.add_paragraph()
    r = p.add_run(
        f'L. Deng, J. Yan, C. Cao, M. Zhao and S. Jin, “{TITLE},” '
        f'in {VENUE}, IEEE, 2026, doi: {DOI}.')
    r.font.size = Pt(10.5)

    # -- provenance -------------------------------------------------------
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run(
        "Bibliographic details and abstract reproduced from the paper's IEEE Xplore record "
        "and attributed to its authors. This is a reference sheet for study purposes; the "
        "full text is available from IEEE Xplore.")
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = GREY

    doc.save(str(OUTPUT))
    print(f"written: {OUTPUT.resolve()}")


if __name__ == "__main__":
    build()
